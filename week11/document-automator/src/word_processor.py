# document-automator/src/word_processor.py
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os
from utils import setup_logging

logger = setup_logging(__name__)

class WordProcessor:
    """
    Handles operations related to Word documents using python-docx.
    """
    def __init__(self):
        logger.info("WordProcessor initialized.")

    def create_document(self):
        """Creates a new blank Word document."""
        document = docx.Document()
        logger.info("New Word document created.")
        return document

    def add_heading(self, document, text, level=1):
        """Adds a heading to the document."""
        document.add_heading(text, level=level)
        logger.info(f"Added heading (Level {level}): '{text}'")

    def add_paragraph(self, document, text, style=None, bold=False, italic=False, font_size=None, align='left'):
        """Adds a paragraph with optional formatting."""
        paragraph = document.add_paragraph(text, style=style)
        
        # Apply run-level formatting if specified
        if paragraph.runs:  # Ensure there's at least one run
            run = paragraph.runs[0]  # Apply to the first run
            if bold:
                run.bold = True
            if italic:
                run.italic = True
            if font_size:
                run.font.size = Pt(font_size)
        
        # Apply paragraph alignment
        if align == 'center':
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == 'right':
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        elif align == 'justify':
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        else:  # Default left
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

        logger.info(f"Added paragraph: '{text[:50]}...'")

    def add_table(self, document, data, include_header=True, header_bold=True, col_widths=None):
        """
        Adds a table to the document.
        `data` should be a list of lists (rows).
        `col_widths` is a list of floats (in inches) for column widths.
        """
        if not data:
            logger.warning("No data provided to add table.")
            return

        rows = len(data)
        cols = len(data[0]) if data else 0
        if rows == 0 or cols == 0:
            logger.warning("Table data is empty or malformed. Skipping table creation.")
            return

        table = document.add_table(rows=rows if include_header else rows, cols=cols)
        table.style = 'Table Grid'  # Apply a basic grid style

        # Set column widths if provided
        if col_widths and len(col_widths) == cols:
            for i, width in enumerate(col_widths):
                table.columns[i].width = Inches(width)

        # Populate table
        for r_idx, row_data in enumerate(data):
            # If header is included in data and we apply special header formatting
            if include_header and r_idx == 0:
                cells = table.rows[0].cells
                for c_idx, cell_data in enumerate(row_data):
                    cells[c_idx].text = str(cell_data)
                    if header_bold:
                        cells[c_idx].paragraphs[0].runs[0].bold = True
            else:
                row_obj = table.rows[r_idx]
                for c_idx, cell_data in enumerate(row_data):
                    row_obj.cells[c_idx].text = str(cell_data)
        
        logger.info(f"Added table with {rows} rows and {cols} columns.")

    def add_picture(self, document, image_path, width_inches=None):
        """Adds an image to the document with optional width scaling."""
        if not os.path.exists(image_path):
            logger.error(f"Image file not found: {image_path}")
            return False
        try:
            if width_inches:
                document.add_picture(image_path, width=Inches(width_inches))
            else:
                document.add_picture(image_path)
            logger.info(f"Added picture: {os.path.basename(image_path)}")
            return True
        except Exception as e:
            logger.error(f"Error adding picture '{image_path}': {e}")
            return False

    def save_document(self, document, output_path):
        """Saves the Word document to a specified path."""
        try:
            document.save(output_path)
            logger.info(f"Word document saved to '{output_path}'.")
            return True
        except Exception as e:
            logger.error(f"Error saving Word document to '{output_path}': {e}")
            return False